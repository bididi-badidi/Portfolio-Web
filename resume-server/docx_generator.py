import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import nsmap

FONT_FAMILY = "Times New Roman"
CONTENT_PT = 10
SUMMARY_PT = 8.5
HEADER_PT = 14
NAVY_BLUE = RGBColor(0x1F, 0x4E, 0x79)


def _add_border_bottom(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _set_spacing(paragraph, before_pt=0, after_pt=0, line_rule=None):
    pf = paragraph.paragraph_format
    if before_pt:
        pf.space_before = Pt(before_pt)
    if after_pt:
        pf.space_after = Pt(after_pt)


def _parse_rich_text(paragraph, text: str, size_pt: float = CONTENT_PT, bold_override: bool = False):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
            run.bold = bold_override
        run.font.name = FONT_FAMILY
        run.font.size = Pt(size_pt)


def _section_title(doc: Document, title: str):
    p = doc.add_paragraph()
    _set_spacing(p, before_pt=4, after_pt=1)
    _add_border_bottom(p)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.name = FONT_FAMILY
    run.font.size = Pt(CONTENT_PT)
    run.font.color.rgb = NAVY_BLUE
    return p


def _subheading(doc: Document, left: str, right: str):
    p = doc.add_paragraph()
    _set_spacing(p, after_pt=0.5)
    run_left = p.add_run(left)
    run_left.bold = True
    run_left.font.name = FONT_FAMILY
    run_left.font.size = Pt(CONTENT_PT)
    run_right = p.add_run(f"\t{right}")
    run_right.font.name = FONT_FAMILY
    run_right.font.size = Pt(CONTENT_PT)
    # right-align the tab stop at ~7.3 inches
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(7.3 * 1440)))
    tabs.append(tab)
    pPr.append(tabs)
    return p


def _bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    _set_spacing(p, after_pt=0.5)
    _parse_rich_text(p, text)
    return p


def _section(doc: Document, title: str, entries: list):
    if not entries:
        return
    _section_title(doc, title)
    for entry in entries:
        header = f"{entry['title']} | {entry['role']}" if entry.get("role") else entry["title"]
        _subheading(doc, header, entry["date"])
        for bullet in entry.get("bullets", []):
            _bullet(doc, bullet)
        doc.add_paragraph().paragraph_format.space_after = Pt(3)


def generate_docx(data: dict) -> bytes:
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.3)
        section.bottom_margin = Inches(0.3)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Remove default paragraph spacing
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)

    header = data["header"]

    # Name + contact
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing(name_p, after_pt=0.5)
    run = name_p.add_run(header["name"])
    run.bold = True
    run.font.name = FONT_FAMILY
    run.font.size = Pt(HEADER_PT)
    if header.get("contact"):
        run2 = name_p.add_run(f" | {header['contact']}")
        run2.font.name = FONT_FAMILY
        run2.font.size = Pt(HEADER_PT)

    # Links line
    links = header.get("links", [])
    if links:
        links_p = doc.add_paragraph()
        links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_spacing(links_p, after_pt=4)
        for i, link in enumerate(links):
            label_run = links_p.add_run(f"{link['label']}: ")
            label_run.bold = True
            label_run.font.name = FONT_FAMILY
            label_run.font.size = Pt(CONTENT_PT)

            # Hyperlink
            r_id = links_p.part.relate_to(link["url"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), r_id)
            r = OxmlElement("w:r")
            rPr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "0000FF")
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rPr.append(color)
            rPr.append(u)
            t = OxmlElement("w:t")
            t.text = link["text"]
            r.append(rPr)
            r.append(t)
            hyperlink.append(r)
            links_p._p.append(hyperlink)

            if i < len(links) - 1:
                sep = links_p.add_run(" | ")
                sep.font.name = FONT_FAMILY
                sep.font.size = Pt(CONTENT_PT)

    # Summary
    _section_title(doc, "Summary")
    summary_p = doc.add_paragraph()
    _set_spacing(summary_p, after_pt=0.5)
    _parse_rich_text(summary_p, data["summary"], size_pt=SUMMARY_PT)

    # Education
    _section_title(doc, "Education")
    for edu in data.get("education", []):
        _subheading(doc, edu["institution"], edu["date"])
        deg_p = doc.add_paragraph()
        _set_spacing(deg_p, after_pt=0.5)
        run = deg_p.add_run(edu["degree"])
        run.italic = True
        run.font.name = FONT_FAMILY
        run.font.size = Pt(CONTENT_PT)
        gpa_p = doc.add_paragraph()
        _set_spacing(gpa_p, after_pt=0.5)
        run = gpa_p.add_run(edu["gpa"])
        run.font.name = FONT_FAMILY
        run.font.size = Pt(CONTENT_PT)

    # Skills
    skills = data.get("skills", {})
    if skills:
        _section_title(doc, "Technical Skills")
        for key, value in skills.items():
            p = doc.add_paragraph()
            _set_spacing(p, after_pt=0.5)
            label = p.add_run(f"{key}: ")
            label.bold = True
            label.font.name = FONT_FAMILY
            label.font.size = Pt(CONTENT_PT)
            val = p.add_run(value)
            val.font.name = FONT_FAMILY
            val.font.size = Pt(CONTENT_PT)

    # Experience sections
    _section(doc, "Work Experiences & Internships", data.get("work_experiences", []))
    _section(doc, "Selected Projects", data.get("personal_projects", []))
    _section(doc, "Leadership Experiences", data.get("leadership_experiences", []))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
